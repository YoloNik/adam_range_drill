"""Excel parsing for question banks.

Expected layout (first worksheet):
    A: question ID
    B: question text
    C, D, E: answer options
    F: correct answer

Column F accepts:
    * the exact text of the correct option (case-insensitive),
    * a letter A/B/C (1st/2nd/3rd option),
    * a column letter C/D/E (detected automatically when any row uses D or E),
    * a number 1/2/3.
A header row is detected and skipped automatically.
"""
import hashlib
import json
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_QUESTIONS = 2000


class ParseError(Exception):
    """Raised when the uploaded workbook cannot be turned into a question bank.

    `key`/`params` and every item of `details` ((key, params) tuples) are i18n keys,
    so the web layer can show the message in the user's language.
    """

    def __init__(self, key, details=None, **params):
        super().__init__(key)
        self.key = key
        self.params = params
        self.details = details or []


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _resolve_answer(answer, options, column_scheme):
    """Return the 0-based index of the correct option, or None if it can't be matched."""
    if not answer:
        return None
    # 1. Exact option text wins
    folded = answer.casefold()
    for i, option in enumerate(options):
        if option and option.casefold() == folded:
            return i
    # 2. Letters / numbers, tolerating "a)", "B." etc.
    key = answer.upper().rstrip(").:").strip()
    letters = {"C": 0, "D": 1, "E": 2} if column_scheme else {"A": 0, "B": 1, "C": 2}
    if key in letters:
        return letters[key]
    if key in {"1", "2", "3"}:
        return int(key) - 1
    return None


def parse_workbook(stream):
    """Parse an uploaded .xlsx stream into a list of question dicts."""
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many different exception types
        raise ParseError("parse.open_failed") from exc

    try:
        ws = wb.worksheets[0]
        rows = []
        for row_number, row in enumerate(ws.iter_rows(min_row=1, max_col=6, values_only=True), start=1):
            cells = [_clean(c) for c in (list(row) + [None] * 6)[:6]]
            if any(cells):
                rows.append((row_number, cells))
            if len(rows) > MAX_QUESTIONS + 1:
                raise ParseError("parse.too_many", n=MAX_QUESTIONS)
    finally:
        wb.close()

    if not rows:
        raise ParseError("parse.empty_sheet")

    # If any answer is "D" or "E", the file uses column letters (C/D/E) for answers
    column_scheme = any(cells[5].upper() in {"D", "E"} for _, cells in rows)

    questions, problems = [], []
    for i, (row_number, cells) in enumerate(rows):
        qid, text, opt1, opt2, opt3, answer = cells
        options = [opt1, opt2, opt3]
        correct = _resolve_answer(answer, options, column_scheme)

        if correct is None:
            if i == 0:
                continue  # first row that doesn't resolve is treated as a header
            problems.append(("parse.row_bad_answer", {"row": row_number, "answer": answer}))
            continue
        if not text:
            problems.append(("parse.row_no_text", {"row": row_number}))
            continue
        if sum(1 for o in options if o) < 2:
            problems.append(("parse.row_few_options", {"row": row_number}))
            continue
        if not options[correct]:
            problems.append(("parse.row_empty_correct", {"row": row_number}))
            continue

        questions.append({
            "external_id": qid or str(row_number),
            "text": text,
            "options": options,
            "correct": correct,
        })

    if problems:
        raise ParseError("parse.problems", details=problems[:15], n=len(problems))
    if not questions:
        raise ParseError("parse.no_questions")
    return questions


SAMPLE_ROWS = {
    "en": [
        ("1", "How should you treat every firearm?", "As if it were always loaded", "As loaded only when the magazine is in", "As unloaded until proven otherwise", "A"),
        ("2", "Where should the muzzle point at all times on the range?", "Toward the instructor", "In a safe direction, downrange", "At the ceiling", "B"),
        ("3", "When may your finger be on the trigger?", "Whenever you hold the weapon", "While walking to the firing line", "Only when sights are on target and you have decided to fire", "C"),
        ("4", "What does the command “CEASE FIRE” require?", "Finish the current magazine", "Stop shooting immediately", "Reload and wait", "B"),
        ("5", "Which protective equipment is mandatory on the firing line?", "Hearing and eye protection", "Gloves only", "None if the range is outdoors", "A"),
        ("6", "What should you do after a misfire?", "Look into the barrel right away", "Keep the weapon pointed downrange and follow the misfire procedure", "Put the weapon down and leave the stand", "B"),
        ("7", "Who gives commands on the range?", "Any experienced shooter", "The range officer", "The shooter on the next lane", "B"),
        ("8", "How do you carry a firearm between stands?", "Unloaded, action open, muzzle in a safe direction", "Loaded with safety on", "Holstered and loaded", "A"),
        ("9", "When are you allowed to go forward of the firing line?", "When you are out of ammunition", "Only after the range officer declares the range safe", "Whenever your lane is empty", "B"),
        ("10", "What must you be sure of before firing?", "Your target and what is beyond it", "Only the distance to the target", "The weather forecast", "A"),
    ],
    "pl": [
        ("1", "Jak należy traktować każdą broń?", "Zawsze jak załadowaną", "Jak załadowaną tylko z włożonym magazynkiem", "Jak rozładowaną, dopóki nie sprawdzisz", "A"),
        ("2", "W którą stronę zawsze powinna być skierowana lufa na strzelnicy?", "W stronę instruktora", "W bezpiecznym kierunku — w stronę kulochwytu", "W sufit", "B"),
        ("3", "Kiedy palec może znajdować się na spuście?", "Zawsze, gdy trzymasz broń", "Podczas przechodzenia na linię ognia", "Tylko gdy przyrządy celownicze są na celu i podjąłeś decyzję o strzale", "C"),
        ("4", "Co oznacza komenda „STOP — PRZERWIJ OGIEŃ”?", "Dokończ strzelanie z magazynka", "Natychmiast przestań strzelać", "Przeładuj i czekaj", "B"),
        ("5", "Jakie środki ochrony są obowiązkowe na linii ognia?", "Ochronniki słuchu i okulary ochronne", "Tylko rękawice", "Żadne, jeśli strzelnica jest otwarta", "A"),
        ("6", "Co zrobić po niewypale?", "Od razu zajrzeć do lufy", "Trzymać broń skierowaną w stronę kulochwytu i postępować zgodnie z procedurą", "Odłożyć broń i odejść ze stanowiska", "B"),
        ("7", "Kto wydaje komendy na strzelnicy?", "Każdy doświadczony strzelec", "Prowadzący strzelanie", "Strzelec z sąsiedniego stanowiska", "B"),
        ("8", "Jak przenosić broń między stanowiskami?", "Rozładowaną, z otwartym zamkiem, lufą w bezpiecznym kierunku", "Załadowaną i zabezpieczoną", "W kaburze, załadowaną", "A"),
        ("9", "Kiedy wolno wyjść przed linię ognia?", "Gdy skończy się amunicja", "Tylko gdy prowadzący strzelanie ogłosi, że jest bezpiecznie", "Gdy twoje stanowisko jest puste", "B"),
        ("10", "Czego musisz być pewien przed oddaniem strzału?", "Celu i tego, co znajduje się za nim", "Tylko odległości do celu", "Prognozy pogody", "A"),
    ],
}

TEMPLATE_HEADERS = {
    "en": ["ID", "Question", "Option A", "Option B", "Option C", "Correct (A/B/C or exact text)"],
    "pl": ["ID", "Pytanie", "Odpowiedź A", "Odpowiedź B", "Odpowiedź C", "Poprawna (A/B/C lub dokładny tekst)"],
}


def content_hash(questions):
    """Fingerprint of a question bank's content, used to detect duplicate uploads.

    Based on what the questions say, not on the file bytes, so re-saving the same Excel file,
    renaming it, reordering rows or answer options, changing IDs / letter case / extra spaces still counts
    as the same bank. `questions` are dicts with "text", "options" and "correct".
    """
    def norm(value):
        return " ".join(str(value).split()).casefold()

    items = sorted(
        json.dumps([norm(q["text"]), sorted(norm(o) for o in q["options"] if o), norm(q["options"][q["correct"]])],
                   ensure_ascii=False)
        for q in questions
    )
    return hashlib.sha256("\n".join(items).encode("utf-8")).hexdigest()


def build_template(lang="pl", sheet_title="Questions"):
    """Create a sample workbook users can download as a starting point."""
    lang = lang if lang in SAMPLE_ROWS else "pl"
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(TEMPLATE_HEADERS[lang])
    for row in SAMPLE_ROWS[lang]:
        ws.append(list(row))

    head_fill = PatternFill("solid", fgColor="344027")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="F4F1E8")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    widths = {"A": 8, "B": 60, "C": 38, "D": 38, "E": 38, "F": 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
